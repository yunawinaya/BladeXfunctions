#!/usr/bin/env python3
"""Emit Item/ItemBulkImportTemplate.xlsx — the upload template for ITEM_BULK_IMPORT.

Columns and allowed values are kept in step with code_parse/code_build in
scratchpad/gen_item_bulk_import_workflow.py. Regenerate with:

    scratchpad/.venv/bin/python3 scratchpad/gen_item_import_template.py
"""
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "Item", "ItemBulkImportTemplate.xlsx")

MAX_ROWS = 500          # mirrors MAX_ROWS in code_parse
PROPERTIES = ["Default", "Product", "Raw Material", "Packaging Material",
              "Work in Progress", "Semi-Finished Goods", "Auxiliary Material"]
BATCH_RULES = ["According To System Settings", "Manual Input"]

# header, required, width, dropdown values, note
COLUMNS = [
    ("Item Code", False, 18, None,
     "Optional. Leave blank and the Items numbering rule assigns the code.\n"
     "If filled, it is used as-is and must not already exist.\n"
     "Example: ITM-0001"),
    ("Item Name", True, 30, None,
     "REQUIRED. Free text.\nExample: Stainless Steel Bolt M8"),
    ("Description", False, 34, None,
     "Optional. Free text.\nExample: M8 x 40mm hex bolt, A2-70"),
    ("Item Category", True, 22, None,
     "REQUIRED. Must match an existing Item Category name exactly.\n"
     "Also decides the item's Costing Method, so there is no Costing Method column.\n"
     "Example: Raw Materials"),
    ("Base UOM", True, 14, None,
     "REQUIRED. Must match an existing UOM name.\n"
     "The UOM conversion row is generated from this.\nExample: PCS"),
    ("Item Properties", False, 22, PROPERTIES,
     "Optional. Blank means Default.\n"
     "Also decides Business Scope, so there is no Business Scope column.\n"
     "Allowed: " + ", ".join(PROPERTIES)),
    ("Item Type", False, 12, ["Goods", "Services"],
     "Optional. Blank means Goods.\nAllowed: Goods, Services"),
    ("Stock Control", False, 14, ["Yes", "No"],
     "Optional. Blank means Yes.\nAllowed: Yes, No"),
    ("Active", False, 10, ["Yes", "No"],
     "Optional. Blank means Yes.\nAllowed: Yes, No"),
    ("Barcode", False, 18, None,
     "Optional. Free text.\nExample: 9551234567890"),
    ("Item Group", False, 18, None,
     "Optional. Must match an existing Item Group code.\nExample: FASTENERS"),
    ("Batch Management", False, 18, ["Yes", "No"],
     "Optional. Blank means No.\nAllowed: Yes, No"),
    ("Batch Number Generation", False, 26, BATCH_RULES,
     "Only when Batch Management is Yes; blank then means "
     "\"According To System Settings\".\n"
     "Leave empty when Batch Management is No.\n"
     "Allowed: " + ", ".join(BATCH_RULES)),
]

FONT = "Arial"
BRAND = "335B57"        # the Item list page's colorPrimary

wb = Workbook()
ws = wb.active
ws.title = "Items"

head_fill = PatternFill("solid", fgColor=BRAND)
req_fill = PatternFill("solid", fgColor="1F3F3C")   # darker, for required columns
thin = Side(style="thin", color="BFBFBF")

for i, (header, required, width, choices, note) in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=i, value=header + (" *" if required else ""))
    cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
    cell.fill = req_fill if required else head_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=thin)
    # The guidance lives in comments, never in a data row -- see the note below.
    cell.comment = Comment(note, "Item Import", height=150, width=330)
    ws.column_dimensions[cell.column_letter].width = width

    if choices:
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(choices) + '"',
            allow_blank=True,
            showDropDown=False,   # openpyxl inverts this: False == show the arrow
        )
        dv.error = "Pick one of: " + ", ".join(choices)
        dv.errorTitle = "Not an allowed value"
        ws.add_data_validation(dv)
        dv.add("%s2:%s%d" % (cell.column_letter, cell.column_letter, MAX_ROWS + 1))

# Set the font on the Normal style rather than per cell: touching each body cell
# would materialise MAX_ROWS empty rows, and the parser would ship them.
wb._named_styles["Normal"].font = Font(name=FONT, size=11)

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:%s1" % ws.cell(row=1, column=len(COLUMNS)).column_letter

wb.save(OUT)
print("wrote %s (%d columns, dropdowns to row %d)" % (OUT, len(COLUMNS), MAX_ROWS + 1))
